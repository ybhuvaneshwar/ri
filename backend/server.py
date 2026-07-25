"""Namma Kavacha — Backend API
FastAPI + MongoDB + JWT auth + Kavacha AI (Claude Sonnet 4.5 via Emergent LLM key)
"""
from fastapi import FastAPI, APIRouter, HTTPException, Depends, UploadFile, File
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Dict, Any
from datetime import datetime, timezone, timedelta
from pathlib import Path
import os, uuid, jwt, bcrypt, logging, io, random, json, asyncio
from collections import Counter, defaultdict

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALGO = "HS256"
EMERGENT_LLM_KEY = os.environ.get('EMERGENT_LLM_KEY', '')

app = FastAPI(title="Namma Kavacha API")
api = APIRouter(prefix="/api")
security = HTTPBearer(auto_error=False)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger("kavacha")

# ---------- MODELS ----------
class LoginRequest(BaseModel):
    email: EmailStr
    password: str

class User(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    email: EmailStr
    name: str
    role: str  # admin | analyst | supervisor
    department: str = "Karnataka State Police"
    designation: str = ""
    unit: str = ""
    district: str = ""
    phone: str = ""
    active: bool = True
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    last_login: Optional[str] = None
    avatar: Optional[str] = None

class UserCreate(BaseModel):
    email: EmailStr
    name: str
    role: str
    password: str
    designation: Optional[str] = ""
    unit: Optional[str] = ""
    district: Optional[str] = ""
    phone: Optional[str] = ""

class Case(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    fir_no: str
    title: str
    crime_head: str
    crime_sub_head: str = ""
    district: str
    unit: str
    status: str = "Open"  # Open, Under Investigation, Chargesheeted, Closed
    severity: str = "Medium"  # Low, Medium, High, Critical
    date_registered: str
    location: str
    lat: float
    lng: float
    description: str = ""
    victims: List[Dict[str, Any]] = []
    accused: List[Dict[str, Any]] = []
    complainant: Dict[str, Any] = {}
    arrests: List[Dict[str, Any]] = []
    chargesheet: Optional[Dict[str, Any]] = None
    io_officer: str = ""
    court: str = ""
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class AuditLog(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    user_email: str
    user_role: str
    action: str
    resource: str
    resource_id: Optional[str] = None
    details: str = ""
    timestamp: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class ChatMessage(BaseModel):
    message: str
    session_id: Optional[str] = None
    context_case_id: Optional[str] = None

# ---------- AUTH HELPERS ----------
def hash_pw(pw: str) -> str:
    return bcrypt.hashpw(pw.encode(), bcrypt.gensalt()).decode()

def verify_pw(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode(), hashed.encode())
    except Exception:
        return False

def make_token(user: dict) -> str:
    payload = {"sub": user["id"], "email": user["email"], "role": user["role"],
               "exp": datetime.now(timezone.utc) + timedelta(days=7)}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGO)

async def current_user(cred: HTTPAuthorizationCredentials = Depends(security)) -> dict:
    if not cred:
        raise HTTPException(401, "Not authenticated")
    try:
        payload = jwt.decode(cred.credentials, JWT_SECRET, algorithms=[JWT_ALGO])
    except jwt.PyJWTError:
        raise HTTPException(401, "Invalid or expired token")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "password": 0})
    if not user or not user.get("active"):
        raise HTTPException(401, "User inactive")
    return user

def require_role(*roles):
    async def dep(user: dict = Depends(current_user)):
        if user["role"] not in roles:
            raise HTTPException(403, f"Requires role in {roles}")
        return user
    return dep

async def log_audit(user: dict, action: str, resource: str, resource_id: str = None, details: str = ""):
    entry = AuditLog(user_email=user["email"], user_role=user["role"], action=action,
                     resource=resource, resource_id=resource_id, details=details)
    await db.audit_logs.insert_one(entry.model_dump())

# ---------- SEED ----------
DISTRICTS = [
    {"name": "Bengaluru City", "lat": 12.9716, "lng": 77.5946},
    {"name": "Mysuru", "lat": 12.2958, "lng": 76.6394},
    {"name": "Mangaluru", "lat": 12.9141, "lng": 74.8560},
    {"name": "Hubballi-Dharwad", "lat": 15.3647, "lng": 75.1240},
    {"name": "Belagavi", "lat": 15.8497, "lng": 74.4977},
    {"name": "Kalaburagi", "lat": 17.3297, "lng": 76.8343},
    {"name": "Tumakuru", "lat": 13.3379, "lng": 77.1173},
    {"name": "Shivamogga", "lat": 13.9299, "lng": 75.5681},
]
CRIME_HEADS = [
    ("Property Crime", "Theft"), ("Property Crime", "Burglary"), ("Property Crime", "Robbery"),
    ("Violent Crime", "Assault"), ("Violent Crime", "Murder"), ("Violent Crime", "Kidnapping"),
    ("Cyber Crime", "Online Fraud"), ("Cyber Crime", "Identity Theft"), ("Cyber Crime", "Phishing"),
    ("Economic Offence", "Cheating"), ("Economic Offence", "Forgery"),
    ("Narcotics", "Drug Peddling"), ("Public Order", "Rioting"),
]
FIRST_NAMES = ["Arjun","Vikram","Ravi","Suresh","Anitha","Priya","Kiran","Rajesh","Manoj","Deepa","Sunita","Naveen","Rahul","Sanjay","Pooja","Meera","Vinod","Ashwin","Lakshmi","Ganesh"]
LAST_NAMES = ["Kumar","Gowda","Rao","Reddy","Shetty","Patil","Iyer","Naidu","Prasad","Bhat","Hegde","Kulkarni","Menon","Nair","Sharma"]

def rand_name():
    return f"{random.choice(FIRST_NAMES)} {random.choice(LAST_NAMES)}"

async def seed_users():
    demos = [
        {"email":"admin@nammakavacha.ai","name":"Arjun Kumar","role":"admin","password":"Admin@123","designation":"System Administrator","unit":"HQ","district":"Bengaluru City","phone":"+91-9876543210"},
        {"email":"analyst@nammakavacha.ai","name":"Priya Shetty","role":"analyst","password":"Analyst@123","designation":"Crime Analyst","unit":"CID","district":"Bengaluru City","phone":"+91-9876543211"},
        {"email":"supervisor@nammakavacha.ai","name":"Vikram Rao","role":"supervisor","password":"Supervisor@123","designation":"Deputy Commissioner","unit":"HQ","district":"Bengaluru City","phone":"+91-9876543212"},
    ]
    for d in demos:
        existing = await db.users.find_one({"email": d["email"]})
        pw_hash = hash_pw(d.pop("password"))
        if not existing:
            u = User(**d).model_dump()
            u["password"] = pw_hash
            await db.users.insert_one(u)
            logger.info(f"Seeded user {d['email']}")
        else:
            # Always reset password to keep demo idempotent
            await db.users.update_one({"email": d["email"]}, {"$set": {"password": pw_hash, "active": True}})

async def seed_cases():
    count = await db.cases.count_documents({})
    if count >= 50:
        return
    await db.cases.delete_many({})
    random.seed(42)
    statuses = ["Open","Under Investigation","Chargesheeted","Closed"]
    severities = ["Low","Medium","High","Critical"]
    cases = []
    now = datetime.now(timezone.utc)
    # Build a network of shared accused for graph analysis
    accused_pool = [{"name": rand_name(), "age": random.randint(19,55), "gender": random.choice(["M","F"])} for _ in range(30)]
    for i in range(80):
        d = random.choice(DISTRICTS)
        ch, csh = random.choice(CRIME_HEADS)
        date = now - timedelta(days=random.randint(1, 365))
        num_accused = random.randint(1,3)
        case_accused = random.sample(accused_pool, num_accused)
        c = Case(
            fir_no=f"FIR/{date.year}/{d['name'][:3].upper()}/{1000+i}",
            title=f"{csh} at {d['name']}",
            crime_head=ch, crime_sub_head=csh,
            district=d["name"], unit=f"{d['name']} Central PS",
            status=random.choices(statuses, weights=[30,35,20,15])[0],
            severity=random.choices(severities, weights=[20,40,30,10])[0],
            date_registered=date.date().isoformat(),
            location=f"Sector {random.randint(1,20)}, {d['name']}",
            lat=d["lat"] + random.uniform(-0.15, 0.15),
            lng=d["lng"] + random.uniform(-0.15, 0.15),
            description=f"A case of {csh.lower()} was reported in {d['name']}. Investigation ongoing with forensic and witness leads being pursued.",
            victims=[{"name": rand_name(), "age": random.randint(18,70), "gender": random.choice(["M","F"]), "contact": f"+91-98{random.randint(10000000,99999999)}"} for _ in range(random.randint(1,2))],
            accused=[{**a, "status": random.choice(["Arrested","Absconding","On Bail"])} for a in case_accused],
            complainant={"name": rand_name(), "phone": f"+91-98{random.randint(10000000,99999999)}", "address": f"House {random.randint(1,200)}, {d['name']}"},
            arrests=[{"person": a["name"], "date": (date + timedelta(days=random.randint(1,20))).date().isoformat(), "officer": rand_name()} for a in case_accused if random.random() > 0.4],
            chargesheet=({"filed_date": (date + timedelta(days=random.randint(30,90))).date().isoformat(), "sections": ["IPC 379","IPC 411"]} if random.random()>0.5 else None),
            io_officer=f"Inspector {rand_name()}",
            court=f"JMFC Court, {d['name']}",
        )
        cases.append(c.model_dump())
    await db.cases.insert_many(cases)
    logger.info(f"Seeded {len(cases)} cases")

# ---------- ROUTES: AUTH ----------
@api.post("/auth/login")
async def login(req: LoginRequest):
    u = await db.users.find_one({"email": req.email})
    if not u or not verify_pw(req.password, u.get("password","")):
        raise HTTPException(401, "Invalid credentials")
    if not u.get("active", True):
        raise HTTPException(403, "Account deactivated")
    await db.users.update_one({"id": u["id"]}, {"$set": {"last_login": datetime.now(timezone.utc).isoformat()}})
    token = make_token(u)
    u.pop("password", None); u.pop("_id", None)
    await log_audit(u, "login", "auth")
    return {"token": token, "user": u}

@api.get("/auth/me")
async def me(user: dict = Depends(current_user)):
    return user

# ---------- ROUTES: USERS (admin) ----------
@api.get("/users")
async def list_users(user: dict = Depends(require_role("admin"))):
    users = await db.users.find({}, {"_id":0, "password":0}).to_list(500)
    return users

@api.post("/users")
async def create_user(payload: UserCreate, user: dict = Depends(require_role("admin"))):
    if await db.users.find_one({"email": payload.email}):
        raise HTTPException(400, "Email already exists")
    u = User(email=payload.email, name=payload.name, role=payload.role,
             designation=payload.designation or "", unit=payload.unit or "",
             district=payload.district or "", phone=payload.phone or "").model_dump()
    u["password"] = hash_pw(payload.password)
    await db.users.insert_one(u)
    await log_audit(user, "create", "user", u["id"], f"Created {payload.email}")
    u.pop("password", None); u.pop("_id", None); return u

@api.delete("/users/{uid}")
async def delete_user(uid: str, user: dict = Depends(require_role("admin"))):
    if uid == user["id"]:
        raise HTTPException(400, "Cannot delete self")
    r = await db.users.delete_one({"id": uid})
    if not r.deleted_count: raise HTTPException(404, "Not found")
    await log_audit(user, "delete", "user", uid)
    return {"ok": True}

@api.patch("/users/{uid}")
async def update_user(uid: str, patch: dict, user: dict = Depends(require_role("admin"))):
    patch.pop("password", None); patch.pop("id", None); patch.pop("_id", None)
    await db.users.update_one({"id": uid}, {"$set": patch})
    await log_audit(user, "update", "user", uid)
    updated = await db.users.find_one({"id": uid}, {"_id":0, "password":0})
    return updated

# ---------- ROUTES: CASES ----------
@api.get("/cases")
async def list_cases(q: Optional[str] = None, district: Optional[str] = None,
                     status: Optional[str] = None, crime_head: Optional[str] = None,
                     limit: int = 100, user: dict = Depends(current_user)):
    query: Dict[str, Any] = {}
    if district: query["district"] = district
    if status: query["status"] = status
    if crime_head: query["crime_head"] = crime_head
    if q:
        query["$or"] = [
            {"fir_no": {"$regex": q, "$options": "i"}},
            {"title": {"$regex": q, "$options": "i"}},
            {"description": {"$regex": q, "$options": "i"}},
            {"location": {"$regex": q, "$options": "i"}},
            {"accused.name": {"$regex": q, "$options": "i"}},
            {"victims.name": {"$regex": q, "$options": "i"}},
        ]
    cases = await db.cases.find(query, {"_id":0}).sort("date_registered", -1).to_list(limit)
    return cases

@api.get("/cases/{cid}")
async def get_case(cid: str, user: dict = Depends(current_user)):
    c = await db.cases.find_one({"id": cid}, {"_id":0})
    if not c: raise HTTPException(404, "Not found")
    return c

@api.post("/cases")
async def create_case(payload: dict, user: dict = Depends(require_role("admin","analyst"))):
    payload.pop("id", None)
    payload.setdefault("date_registered", datetime.now(timezone.utc).date().isoformat())
    c = Case(**payload).model_dump()
    await db.cases.insert_one(c)
    await log_audit(user, "create", "case", c["id"], c["fir_no"])
    c.pop("_id", None)
    return c

@api.patch("/cases/{cid}")
async def update_case(cid: str, patch: dict, user: dict = Depends(require_role("admin","analyst"))):
    patch.pop("id", None); patch.pop("_id", None)
    r = await db.cases.update_one({"id": cid}, {"$set": patch})
    if not r.matched_count: raise HTTPException(404, "Not found")
    await log_audit(user, "update", "case", cid)
    return await db.cases.find_one({"id": cid}, {"_id":0})

@api.delete("/cases/{cid}")
async def delete_case(cid: str, user: dict = Depends(require_role("admin"))):
    r = await db.cases.delete_one({"id": cid})
    if not r.deleted_count: raise HTTPException(404, "Not found")
    await log_audit(user, "delete", "case", cid)
    return {"ok": True}

# ---------- ROUTES: ANALYTICS / DASHBOARD ----------
@api.get("/dashboard/kpis")
async def kpis(user: dict = Depends(current_user)):
    total = await db.cases.count_documents({})
    open_c = await db.cases.count_documents({"status": {"$in": ["Open","Under Investigation"]}})
    closed = await db.cases.count_documents({"status": "Closed"})
    critical = await db.cases.count_documents({"severity": "Critical"})
    chargesheeted = await db.cases.count_documents({"status": "Chargesheeted"})
    # Recent 30d vs previous 30d
    now = datetime.now(timezone.utc)
    d30 = (now - timedelta(days=30)).date().isoformat()
    d60 = (now - timedelta(days=60)).date().isoformat()
    recent = await db.cases.count_documents({"date_registered": {"$gte": d30}})
    prev = await db.cases.count_documents({"date_registered": {"$gte": d60, "$lt": d30}})
    trend = ((recent - prev) / prev * 100) if prev else 0
    return {"total": total, "open": open_c, "closed": closed, "critical": critical,
            "chargesheeted": chargesheeted, "recent_30d": recent, "trend_pct": round(trend,1)}

@api.get("/analytics/by-district")
async def by_district(user: dict = Depends(current_user)):
    cases = await db.cases.find({}, {"_id":0, "district":1, "severity":1, "status":1}).to_list(5000)
    out: Dict[str, Dict[str,int]] = defaultdict(lambda: {"total":0,"open":0,"critical":0})
    for c in cases:
        out[c["district"]]["total"] += 1
        if c.get("status") in ["Open","Under Investigation"]:
            out[c["district"]]["open"] += 1
        if c.get("severity") == "Critical":
            out[c["district"]]["critical"] += 1
    return [{"district": k, **v} for k,v in out.items()]

@api.get("/analytics/by-category")
async def by_category(user: dict = Depends(current_user)):
    cases = await db.cases.find({}, {"_id":0, "crime_head":1}).to_list(5000)
    c = Counter([x["crime_head"] for x in cases])
    return [{"category": k, "count": v} for k,v in c.most_common()]

@api.get("/analytics/timeline")
async def timeline(days: int = 90, user: dict = Depends(current_user)):
    now = datetime.now(timezone.utc)
    since = (now - timedelta(days=days)).date().isoformat()
    cases = await db.cases.find({"date_registered": {"$gte": since}}, {"_id":0,"date_registered":1,"crime_head":1}).to_list(5000)
    buckets: Dict[str,int] = defaultdict(int)
    for c in cases:
        # bucket by week
        d = datetime.fromisoformat(c["date_registered"])
        wk = (d - timedelta(days=d.weekday())).date().isoformat()
        buckets[wk] += 1
    return [{"week": k, "count": v} for k,v in sorted(buckets.items())]

@api.get("/map/hotspots")
async def hotspots(user: dict = Depends(current_user)):
    cases = await db.cases.find({}, {"_id":0, "lat":1,"lng":1,"district":1,"crime_head":1,"severity":1,"fir_no":1,"id":1,"title":1}).to_list(5000)
    return cases

@api.get("/network/graph")
async def network_graph(user: dict = Depends(current_user)):
    cases = await db.cases.find({}, {"_id":0, "id":1,"fir_no":1,"title":1,"accused":1,"district":1,"severity":1}).to_list(2000)
    nodes: Dict[str, dict] = {}
    links: List[dict] = []
    for c in cases:
        cn = f"case:{c['id']}"
        nodes[cn] = {"id": cn, "name": c["fir_no"], "type":"case", "district": c["district"],
                     "severity": c.get("severity","Medium"), "label": c["title"]}
        for a in c.get("accused", []):
            an = f"person:{a['name']}"
            nodes[an] = {"id": an, "name": a["name"], "type":"person", "label": a["name"]}
            links.append({"source": cn, "target": an})
    return {"nodes": list(nodes.values()), "links": links}

@api.get("/predictions/hotspots")
async def predictions(user: dict = Depends(current_user)):
    # Simple heuristic: districts with rising recent counts + high severity share
    now = datetime.now(timezone.utc)
    d30 = (now - timedelta(days=30)).date().isoformat()
    d90 = (now - timedelta(days=90)).date().isoformat()
    recent = await db.cases.find({"date_registered": {"$gte": d30}}, {"_id":0}).to_list(5000)
    baseline = await db.cases.find({"date_registered": {"$gte": d90, "$lt": d30}}, {"_id":0}).to_list(5000)
    def by_d(rows):
        c = Counter(); crit = Counter()
        for r in rows:
            c[r["district"]] += 1
            if r.get("severity") == "Critical": crit[r["district"]] += 1
        return c, crit
    rc, rcrit = by_d(recent); bc, _ = by_d(baseline)
    preds = []
    for d in DISTRICTS:
        r = rc.get(d["name"],0); b = max(bc.get(d["name"],1)/2, 1)  # avg per 30d
        growth = (r - b) / b
        confidence = min(0.95, 0.55 + max(0,growth)*0.4 + rcrit.get(d["name"],0)*0.05)
        risk = "High" if growth > 0.3 else "Medium" if growth > 0 else "Low"
        preds.append({"district": d["name"], "lat": d["lat"], "lng": d["lng"],
                      "recent_30d": r, "baseline_30d": round(b,1),
                      "growth_pct": round(growth*100,1), "risk": risk,
                      "confidence": round(confidence*100,1),
                      "reason": f"{'Increased' if growth>0 else 'Stable'} case volume ({r} recent vs {round(b,1)} avg), {rcrit.get(d['name'],0)} critical cases."})
    return sorted(preds, key=lambda x: -x["confidence"])

@api.get("/cases/{cid}/similar")
async def similar_cases(cid: str, user: dict = Depends(current_user)):
    target = await db.cases.find_one({"id": cid}, {"_id":0})
    if not target: raise HTTPException(404, "Not found")
    others = await db.cases.find({"id": {"$ne": cid}}, {"_id":0}).to_list(2000)
    scored = []
    for c in others:
        score = 0
        if c["crime_head"] == target["crime_head"]: score += 40
        if c["crime_sub_head"] == target["crime_sub_head"]: score += 25
        if c["district"] == target["district"]: score += 15
        # geo distance
        dx = abs(c["lat"] - target["lat"]); dy = abs(c["lng"] - target["lng"])
        if dx + dy < 0.1: score += 10
        # accused overlap
        ta = {a["name"] for a in target.get("accused",[])}
        ca = {a["name"] for a in c.get("accused",[])}
        if ta & ca: score += 20 * len(ta & ca)
        if score > 40:
            scored.append({**c, "match_score": score})
    scored.sort(key=lambda x: -x["match_score"])
    return scored[:10]

# ---------- KAVACHA AI ----------
async def build_kavacha_context(case_id: Optional[str] = None) -> str:
    """Build a compact context for the AI (case corpus summary + optional focused case)."""
    total = await db.cases.count_documents({})
    by_cat = await by_category({"role":"admin"}) if False else None  # placeholder
    cats = await db.cases.aggregate([{"$group":{"_id":"$crime_head","c":{"$sum":1}}}]).to_list(50)
    districts = await db.cases.aggregate([{"$group":{"_id":"$district","c":{"$sum":1}}}]).to_list(50)
    ctx = f"Namma Kavacha corpus: {total} cases across Karnataka.\n"
    ctx += "Cases by category: " + ", ".join([f"{c['_id']}={c['c']}" for c in cats]) + "\n"
    ctx += "Cases by district: " + ", ".join([f"{d['_id']}={d['c']}" for d in districts]) + "\n"
    if case_id:
        c = await db.cases.find_one({"id": case_id}, {"_id":0})
        if c:
            ctx += f"\n[FOCUSED CASE] FIR {c['fir_no']} — {c['title']} — {c['district']} — Status: {c['status']} — Severity: {c['severity']}\n"
            ctx += f"Description: {c['description']}\n"
            ctx += f"Accused: {[a['name'] for a in c.get('accused',[])]}\n"
            ctx += f"Victims: {[v['name'] for v in c.get('victims',[])]}\n"
    # recent 5 cases as sample
    recent = await db.cases.find({}, {"_id":0}).sort("date_registered",-1).limit(5).to_list(5)
    ctx += "\nRecent cases:\n"
    for c in recent:
        ctx += f"- {c['fir_no']} | {c['title']} | {c['district']} | {c['status']} | {c['date_registered']}\n"
    return ctx

@api.post("/kavacha/chat")
async def kavacha_chat(payload: ChatMessage, user: dict = Depends(current_user)):
    from emergentintegrations.llm.chat import LlmChat, UserMessage, TextDelta, StreamDone
    session_id = payload.session_id or str(uuid.uuid4())
    context = await build_kavacha_context(payload.context_case_id)
    system = (
        "You are Kavacha AI, an intelligence copilot for Karnataka State Police officers. "
        "You help investigators analyze crime patterns, cases, and relationships. "
        "Always cite specific FIR numbers when referencing cases. Be concise, data-driven, "
        "and end responses with a 'Next steps' recommendation when appropriate. "
        "Format citations as [FIR/YYYY/DIS/NNNN]. Use markdown for structure.\n\n"
        f"CURRENT CONTEXT:\n{context}"
    )
    chat = LlmChat(api_key=EMERGENT_LLM_KEY, session_id=session_id, system_message=system).with_model("anthropic","claude-sonnet-4-5-20250929")
    async def gen():
        try:
            async for ev in chat.stream_message(UserMessage(text=payload.message)):
                if isinstance(ev, TextDelta):
                    yield f"data: {json.dumps({'delta': ev.content})}\n\n"
                elif isinstance(ev, StreamDone):
                    yield f"data: {json.dumps({'done': True, 'session_id': session_id})}\n\n"
                    break
        except Exception as e:
            logger.exception("Kavacha stream failed")
            yield f"data: {json.dumps({'error': str(e)})}\n\n"
    await log_audit(user, "chat", "kavacha_ai", session_id, payload.message[:100])
    return StreamingResponse(gen(), media_type="text/event-stream",
                             headers={"Cache-Control":"no-cache","X-Accel-Buffering":"no"})

@api.post("/kavacha/report/{cid}")
async def generate_report(cid: str, user: dict = Depends(current_user)):
    from emergentintegrations.llm.chat import LlmChat, UserMessage
    c = await db.cases.find_one({"id": cid}, {"_id":0})
    if not c: raise HTTPException(404, "Case not found")
    similar = await similar_cases(cid, user)
    system = "You are Kavacha AI, generating a formal investigation summary report for Karnataka State Police."
    prompt = (
        f"Generate a formal investigation summary for FIR {c['fir_no']} — {c['title']}. "
        f"Include sections: 1) Executive Summary, 2) Case Details, 3) Persons Involved, "
        f"4) Investigation Status, 5) Similar/Related Cases, 6) Recommended Next Steps. "
        f"Use markdown. Case data:\n{json.dumps(c, default=str)[:3000]}\n\n"
        f"Similar cases identified: {[s['fir_no'] for s in similar[:5]]}"
    )
    chat = LlmChat(api_key=EMERGENT_LLM_KEY, session_id=f"report-{cid}", system_message=system).with_model("anthropic","claude-sonnet-4-5-20250929")
    try:
        resp = await chat.send_message(UserMessage(text=prompt))
        await log_audit(user, "generate_report", "case", cid)
        return {"case_id": cid, "fir_no": c["fir_no"], "report": resp, "generated_at": datetime.now(timezone.utc).isoformat()}
    except Exception as e:
        raise HTTPException(500, f"AI report generation failed: {e}")

# ---------- UPLOADS ----------
@api.post("/upload/excel")
async def upload_excel(file: UploadFile = File(...), user: dict = Depends(require_role("admin"))):
    import openpyxl
    data = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(data))
    imported = 0; errors: List[str] = []
    if "CaseMaster" in wb.sheetnames:
        ws = wb["CaseMaster"]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            try:
                rec = dict(zip(headers, row))
                if not rec.get("fir_no"): continue
                # coerce lat/lng if provided
                lat = float(rec.get("lat") or 12.97)
                lng = float(rec.get("lng") or 77.59)
                c = Case(fir_no=str(rec["fir_no"]), title=str(rec.get("title","Imported Case")),
                         crime_head=str(rec.get("crime_head","Property Crime")),
                         crime_sub_head=str(rec.get("crime_sub_head","")),
                         district=str(rec.get("district","Bengaluru City")),
                         unit=str(rec.get("unit","")),
                         status=str(rec.get("status","Open")),
                         severity=str(rec.get("severity","Medium")),
                         date_registered=str(rec.get("date_registered", datetime.now(timezone.utc).date().isoformat())),
                         location=str(rec.get("location","")),
                         lat=lat, lng=lng, description=str(rec.get("description","")))
                await db.cases.update_one({"fir_no": c.fir_no}, {"$set": c.model_dump()}, upsert=True)
                imported += 1
            except Exception as e:
                errors.append(str(e))
    await log_audit(user, "upload", "excel", details=f"{imported} records")
    return {"imported": imported, "errors": errors[:5]}

@api.post("/upload/pdf-extract")
async def upload_pdf(file: UploadFile = File(...), user: dict = Depends(require_role("admin","analyst"))):
    import pypdf
    data = await file.read()
    reader = pypdf.PdfReader(io.BytesIO(data))
    text = "\n".join([p.extract_text() or "" for p in reader.pages])[:4000]
    from emergentintegrations.llm.chat import LlmChat, UserMessage
    system = "Extract structured FIR data from Indian police documents. Return strict JSON."
    prompt = (
        "Extract these fields as JSON: fir_no, title, crime_head, crime_sub_head, district, "
        "unit, status, severity, date_registered (YYYY-MM-DD), location, description, "
        "complainant (name/phone), victims (array of name/age/gender), accused (array of name/age/gender). "
        "If a field is unknown, use empty string or empty array. Return ONLY JSON, no prose.\n\n"
        f"PDF TEXT:\n{text}"
    )
    chat = LlmChat(api_key=EMERGENT_LLM_KEY, session_id=f"pdf-{uuid.uuid4()}", system_message=system).with_model("anthropic","claude-sonnet-4-5-20250929")
    try:
        resp = await chat.send_message(UserMessage(text=prompt))
        # crude JSON extraction
        s = resp.find("{"); e = resp.rfind("}")
        extracted = json.loads(resp[s:e+1]) if s >=0 else {}
        return {"extracted": extracted, "raw_preview": text[:500]}
    except Exception as e:
        raise HTTPException(500, f"Extraction failed: {e}")

# ---------- AUDIT ----------
@api.get("/audit")
async def get_audit(limit: int = 200, user: dict = Depends(require_role("admin","supervisor"))):
    logs = await db.audit_logs.find({}, {"_id":0}).sort("timestamp",-1).to_list(limit)
    return logs

# ---------- NOTIFICATIONS ----------
@api.get("/notifications")
async def notifications(user: dict = Depends(current_user)):
    # Derived alerts
    critical = await db.cases.count_documents({"severity":"Critical","status":{"$in":["Open","Under Investigation"]}})
    preds = await predictions(user)
    high_risk = [p for p in preds if p["risk"]=="High"][:3]
    notifs = []
    if critical:
        notifs.append({"id":"n1","severity":"crimson","title":f"{critical} critical cases open","body":"Immediate attention required.","timestamp": datetime.now(timezone.utc).isoformat()})
    for p in high_risk:
        notifs.append({"id":f"n-{p['district']}","severity":"amber",
                       "title":f"Hotspot alert: {p['district']}",
                       "body":f"Predicted {p['growth_pct']}% surge — confidence {p['confidence']}%.",
                       "timestamp": datetime.now(timezone.utc).isoformat()})
    notifs.append({"id":"n-info","severity":"emerald","title":"Kavacha AI is online","body":"Copilot ready to assist.","timestamp": datetime.now(timezone.utc).isoformat()})
    return notifs

# ---------- Health ----------
@api.get("/")
async def root():
    return {"service":"Namma Kavacha","status":"operational","time": datetime.now(timezone.utc).isoformat()}

app.include_router(api)
app.add_middleware(CORSMiddleware, allow_credentials=True,
                   allow_origins=os.environ.get('CORS_ORIGINS','*').split(','),
                   allow_methods=["*"], allow_headers=["*"])

@app.on_event("startup")
async def startup():
    await seed_users()
    await seed_cases()
    logger.info("Namma Kavacha ready.")

@app.on_event("shutdown")
async def shutdown():
    client.close()
